# -*- coding: utf-8 -*-
"""Build the 2026 Yunqiao Cup submission booklet (DOCX) for Saker Energy.
Facts come only from the repository records and live API runs (see outline.md).
"""
import docx
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = r'E:\saker-energy-main\docs\integrations\wanzhi\submission\2026云桥杯-SakerEnergy充储场站证据诊断助手-作品说明书.docx'

BLACK = RGBColor(0, 0, 0)
HEI = '黑体'
SONG = '宋体'
KAI = '楷体'


def set_run_font(run, cn_font, size, bold=False, color=BLACK):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), cn_font)


def para(doc, text, style=None, align=None, indent_chars=None, space_after=None,
         space_before=None, line=None, font=None, size=12, bold=False):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if indent_chars:
        ppr = p._p.get_or_add_pPr()
        ind = ppr.find(qn('w:ind'))
        if ind is None:
            ind = OxmlElement('w:ind')
            ppr.append(ind)
        ind.set(qn('w:firstLineChars'), str(indent_chars * 100))
        ind.set(qn('w:firstLine'), str(int(indent_chars * size * 20)))
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    if space_before is not None:
        p.paragraph_format.space_before = Pt(space_before)
    if line is not None:
        p.paragraph_format.line_spacing = line
    run = p.add_run(text)
    if font:
        set_run_font(run, font, size, bold)
    return p


def heading(doc, text, level):
    h = doc.add_heading('', level=level)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for r in h.runs:
        r.text = ''
    run = h.add_run(text)
    sizes = {1: 16, 2: 14, 3: 12}
    set_run_font(run, HEI, sizes.get(level, 12), bold=True, color=BLACK)
    hpr = h._p.get_or_add_pPr()
    for tag in ('w:keepNext', 'w:keepLines'):
        el = hpr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            hpr.append(el)
    return h


def caption(doc, text):
    p = para(doc, text, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4,
             space_after=8, font=SONG, size=10.5)
    return p


def add_page_field(section, restart=None):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), ' PAGE ')
    p._p.append(fld)
    if restart is not None:
        sectPr = section._sectPr
        pg = sectPr.find(qn('w:pgNumType'))
        if pg is None:
            pg = OxmlElement('w:pgNumType')
            sectPr.append(pg)
        pg.set(qn('w:start'), str(restart))


TOC_ITEMS = [
    ('1 项目概述', 3, 1),
    ('1.1 背景', 3, 2),
    ('1.2 作品定位', 3, 2),
    ('1.3 核心成果', 3, 2),
    ('2 业务场景与问题定义', 3, 1),
    ('2.1 典型场景', 3, 2),
    ('2.2 待解问题', 3, 2),
    ('2.3 诊断合同（防误判约束）', 3, 2),
    ('3 Agent 设计与诊断合同', 4, 1),
    ('3.1 输入证据包', 4, 2),
    ('3.2 决策输出', 4, 2),
    ('3.3 本地校验合同', 4, 2),
    ('4 系统架构与工程化实现', 4, 1),
    ('4.1 调用链', 4, 2),
    ('4.2 本地工具链', 4, 2),
    ('4.3 工程化保障', 4, 2),
    ('5 真实运行证据', 5, 1),
    ('5.1 六轮真实 API 调用', 5, 2),
    ('5.2 补证闭环行为', 5, 2),
    ('6 安全与可信设计', 5, 1),
    ('7 商业价值与应用前景', 6, 1),
    ('8 边界声明', 6, 1),
]


def add_toc(doc):
    # Standard TOC field (begin/instrText/separate + static entries + end)
    # so Word can F9-refresh it; the static entries keep the PDF render usable.
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    r1 = p.add_run()
    fld_b = OxmlElement('w:fldChar')
    fld_b.set(qn('w:fldCharType'), 'begin')
    r1._r.append(fld_b)
    r2 = p.add_run()
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = r' TOC \o "1-2" \h \z \u '
    r2._r.append(instr)
    r3 = p.add_run()
    fld_s = OxmlElement('w:fldChar')
    fld_s.set(qn('w:fldCharType'), 'separate')
    r3._r.append(fld_s)
    for text, page, level in TOC_ITEMS:
        ep = doc.add_paragraph()
        ep.paragraph_format.line_spacing = 1.0
        ep.paragraph_format.space_after = Pt(2)
        ppr = ep._p.get_or_add_pPr()
        tabs = OxmlElement('w:tabs')
        tab = OxmlElement('w:tab')
        tab.set(qn('w:val'), 'right')
        tab.set(qn('w:leader'), 'dot')
        tab.set(qn('w:pos'), str(int(16.0 * 567)))
        tabs.append(tab)
        ppr.append(tabs)
        if level == 2:
            ind = OxmlElement('w:ind')
            ind.set(qn('w:leftChars'), '200')
            ppr.append(ind)
        er1 = ep.add_run(text)
        set_run_font(er1, SONG, 12)
        er2 = ep.add_run('\t' + str(page))
        set_run_font(er2, SONG, 12)
    ep = doc.add_paragraph()
    ep.paragraph_format.line_spacing = 1.0
    r_end = ep.add_run()
    fld_e = OxmlElement('w:fldChar')
    fld_e.set(qn('w:fldCharType'), 'end')
    r_end._r.append(fld_e)


def table(doc, rows, widths=None, header_bg='D9D9D9'):
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = t._tbl.tblPr
    cant = tblPr.find(qn('w:cantSplit'))
    if cant is None:
        cant = OxmlElement('w:cantSplit')
        tblPr.append(cant)
    for ri, row in enumerate(rows):
        tr = t.add_row()._tr
        trPr = tr.get_or_add_trPr()
        if ri == 0:
            header = OxmlElement('w:tblHeader')
            trPr.append(header)
        trPr.append(OxmlElement('w:cantSplit'))
        for ci, value in enumerate(row):
            cell = t.rows[ri].cells[ci]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cell.width = Cm(widths[ci])
            cp = cell.paragraphs[0]
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER if (ri == 0 or ci == 0) else WD_ALIGN_PARAGRAPH.LEFT
            cp.paragraph_format.line_spacing = 1.0
            cp.paragraph_format.space_after = Pt(2)
            cp.paragraph_format.space_before = Pt(2)
            cp.paragraph_format.first_line_indent = Cm(0)
            run = cp.add_run(str(value))
            set_run_font(run, SONG, 10.5, bold=(ri == 0))
            if ri == 0:
                tcPr = cell._tc.get_or_add_tcPr()
                shd = OxmlElement('w:shd')
                shd.set(qn('w:val'), 'clear')
                shd.set(qn('w:fill'), header_bg)
                tcPr.append(shd)
    return t


def main():
    doc = Document()

    # Page setup: A4, 2.5cm margins
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    for attr in ('top_margin', 'bottom_margin', 'left_margin', 'right_margin'):
        setattr(sec, attr, Cm(2.5))

    # Normal style
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(12)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), SONG)

    # ---- Cover ----
    para(doc, '', size=12)
    para(doc, '2026「云桥杯」工业智能体大赛作品说明书', align=WD_ALIGN_PARAGRAPH.CENTER,
         font=HEI, size=18, bold=True, space_before=90, space_after=18)
    para(doc, '—— 充储场站证据诊断助手 ——', align=WD_ALIGN_PARAGRAPH.CENTER,
         font=KAI, size=12, space_after=60)
    for line in ('作品名称：Saker Energy 充储场站证据诊断助手',
                 '参赛赛道：工业智能体应用赛',
                 'Agent ID：d86bc562cb9b450b（已发布）',
                 'Agent 公开链接：https://www.wanzhi.com/studio/agent/d86bc562cb9b450b',
                 '日期：2026 年 10 月 29 日'):
        para(doc, line, align=WD_ALIGN_PARAGRAPH.CENTER, font=SONG, size=12,
             space_after=6)
    doc.add_page_break()

    # ---- TOC ----
    para(doc, '目 录', align=WD_ALIGN_PARAGRAPH.CENTER, font=HEI, size=16,
         bold=True, space_after=12)
    add_toc(doc)
    doc.add_page_break()

    # ---- Body section (restart page numbers) ----
    body = doc.add_section(WD_SECTION.NEW_PAGE)
    for attr in ('top_margin', 'bottom_margin', 'left_margin', 'right_margin'):
        setattr(body, attr, Cm(2.5))
    add_page_field(body, restart=1)

    # 1 项目概述
    heading(doc, '1 项目概述', 1)
    heading(doc, '1.1 背景', 2)
    para(doc, '充储场站（储能与充电一体）在运行中常出现多设备功率异常告警。现场可用的证据分散在遥测、拓扑、调度记录与告警日志中，运维人员难以快速、可信地判定根因候选，误判会直接导致处置方向错误。本作品以「证据驱动、只读查询、人工审批」为原则，构建一个面向充储场站告警证据链诊断的智能体。', indent_chars=2)
    heading(doc, '1.2 作品定位', 2)
    para(doc, '在万智平台发布专属员工「Saker Energy 充储场站证据诊断助手」（Agent ID：d86bc562cb9b450b，version 3），通过万智开放 API（https://www.wanzhi.com/app/api/v1/open-api）接入本地工具链，实现「证据不足不妄断、主动请求只读补证、无法定论转人工」的诊断闭环。', indent_chars=2)
    heading(doc, '1.3 核心成果', 2)
    para(doc, '已完成 6 轮真实 API 调用，全部通过本地 Schema 与语义合同校验（VALIDATED_ADVISORY）；其中一轮完整验证了补证闭环：智能体请求只读补证 → 本地执行模拟遥测查询 → 结果回发同一会话 → 智能体基于新证据修订候选并正确使用反证，最终因预算与证据价值耗尽转人工核验。', indent_chars=2)

    # 2 业务场景与问题定义
    heading(doc, '2 业务场景与问题定义', 1)
    heading(doc, '2.1 典型场景', 2)
    para(doc, '参考拓扑：10kV 进线 → 变压器降压 → 0.4kV 母线 BUS-B，母线下挂 500kW/1000kWh 储能柜（ESS-01）、500kW 储能变流器（PCS-01）、120kW 双枪充电桩（CP-05）及母线电压测点 METER-01（额定 380V，运行范围 ±7%）。', indent_chars=2)
    heading(doc, '2.2 待解问题', 2)
    para(doc, '在某 1 分钟时间窗内，PCS-01 充电功率由 180kW 降至 110kW、CP-05 由 96kW 降至 60kW，两者同属 BUS-B。仅凭两项功率证据无法区分：上游供电受限/母线电压异常、站级调度限功率/峰谷策略、测点时间戳或读数质量异常等候选原因。', indent_chars=2)
    heading(doc, '2.3 诊断合同（防误判约束）', 2)
    para(doc, '智能体输出必须遵守：证据不足时 abstain=true，不做确定性根因结论；候选原因只能引用输入证据；补证请求仅限白名单只读查询并受查询预算约束；缺失数据返回 missing、不编造读数；母线电压正常不能直接证明调度限功率；无法定论时输出人工核验清单。', indent_chars=2)

    # 3 Agent 设计
    heading(doc, '3 Agent 设计与诊断合同', 1)
    heading(doc, '3.1 输入证据包', 2)
    para(doc, '每次诊断输入为一个结构化证据包（JSON），字段包括：schema_version、event_id、input_revision、round_id、trace_id、data_origin、window（时间窗）、evidence（测点/量纲/采样/质量/来源）、missing_information、available_queries（白名单只读查询）与 query_budget_remaining。示例见 evidence-package.example.json 与 industry-demo.input.json。', indent_chars=2)
    heading(doc, '3.2 决策输出', 2)
    para(doc, '智能体输出决策 JSON，字段包括：status（NEEDS_EVIDENCE/NEEDS_HUMAN）、abstain、summary、hypotheses（候选原因及支持/反证证据引用）、requested_evidence（补证请求及理由）、human_checks（人工核验清单）。完整约束见 decision.schema.json。', indent_chars=2)
    heading(doc, '3.3 本地校验合同', 2)
    para(doc, '本地工具链对每个真实输出执行 Schema 校验与语义校验：轮次/事件/数据源与输入一致、证据引用不越界、支持与反证不重叠、补证请求受预算约束、低质量证据必须 abstain、禁止确认物理根因（advisory 模式）。', indent_chars=2)

    # 4 系统架构与工程化实现
    heading(doc, '4 系统架构与工程化实现', 1)
    heading(doc, '4.1 调用链', 2)
    para(doc, '整体链路：Next.js 大屏 → Go 任务 API（规划）→ Python 万智适配器 → 万智开放 API。适配器以 SSE 流式接收模型输出（text / reasoning / message_result / finish_reason），完成事件为 finish_reason；多轮补证在同一会话内进行，不重复创建会话。', indent_chars=2)
    heading(doc, '4.2 本地工具链', 2)
    para(doc, 'backend/wanzhi.py 提供命令行：--configure（密钥加密录入）、--check、--probe（只读探测）、--validate（离线校验）、--run（单轮真实调用）、--resume（只读对账、绝不重发）、--followup（补证闭环第二轮）。密钥经 Windows DPAPI 加密存于 backend/.private（gitignore）；所有输出经 redact 脱敏。', indent_chars=2)
    heading(doc, '4.3 工程化保障', 2)
    para(doc, '本地单元测试 50 项（合同、传输、检查点、凭据）全部通过；失败状态机区分 UNKNOWN_REMOTE_STATE（创建/发送阶段不确定）与 RESULT_NOT_VALIDATED（结果未通过校验），不重发、可对账；历史提交经密钥扫描确认无真实凭据入库。', indent_chars=2)

    # 5 真实运行证据
    heading(doc, '5 真实运行证据', 1)
    heading(doc, '5.1 六轮真实 API 调用', 2)
    caption(doc, '表1 真实 API 调用汇总（全部 VALIDATED_ADVISORY）')
    table(doc, [
        ['轮次', '场景', 'session_id', 'decision', '完成事件'],
        ['1', '合成样例·多设备功率下降', 'd3e07a0294e1430b', 'NEEDS_EVIDENCE·请求Q01', 'finish_reason'],
        ['2', '合成·查询预算耗尽', 'cf66bbd4c90c48c5', 'NEEDS_HUMAN·转人工', 'finish_reason'],
        ['3', '合成·电压400→360V', 'a77b2f735ff44247', 'NEEDS_HUMAN·转人工', 'finish_reason'],
        ['4', '合成·电压缺失', 'b3ed93633ed24106', 'NEEDS_HUMAN·不编造', 'finish_reason'],
        ['5', '行业·BUS-B功率下降', '8f8b1405cefe45c7', 'NEEDS_EVIDENCE·请求Q01', 'finish_reason'],
        ['6', '行业·补证闭环（同会话）', '8f8b1405cefe45c7', 'NEEDS_HUMAN·修订转人工', 'finish_reason'],
    ], widths=[1.4, 4.6, 4.2, 4.2, 2.6])
    para(doc, '观测到的 SSE 单轮帧数：1124（轮1）、1087（轮5）、1008（轮6）。真实鉴权契约：model_api_key 字段必填，取值须为有效模型调用凭据（大赛发放的大模型 API Key）；平台 OpenAPI Key 与 Agent 专属 Key 均被模型侧拒绝，过程与适配已记录在 observations。', indent_chars=2)
    heading(doc, '5.2 补证闭环行为', 2)
    caption(doc, '表2 行业场景两轮行为对比')
    table(doc, [
        ['维度', '首轮（轮5）', '补证第二轮（轮6）'],
        ['输入', 'E01/E02/E03，缺母线电压', '+E04：母线电压396V（本地只读模拟）'],
        ['候选', 'H1上游供电 / H2站级调度 / H3测量对时', 'H1保留但增加 against=[E04]；H2不误升；H3保留'],
        ['补证请求', 'Q01（读METER-01母线电压）', '无：Q01已执行、重复读取无新增价值，转人工'],
        ['结论', 'NEEDS_EVIDENCE', 'NEEDS_HUMAN（7项人工核验）'],
    ], widths=[2.4, 7.0, 7.0])
    para(doc, '第二轮的关键修订行为符合诊断合同：母线电压正常削弱了「电压异常」分支（正确使用反证），但不据此升级「调度限功率」候选；对单点模拟读数主动标注需交叉验证；预算与证据价值评估后不再请求补证，转人工核验。', indent_chars=2)

    # 6 安全与可信设计
    heading(doc, '6 安全与可信设计', 1)
    para(doc, '本作品面向工业现场，安全边界明确：只读查询白名单（telemetry_window 等），不含任何设备控制指令；无法定论时强制转人工核验（human_checks）；证据来源可追溯（data_origin/source/provenance，模拟值显式标注）；密钥仅存服务端加密配置，不入前端与公开仓库。', indent_chars=2)

    # 7 商业价值与前景
    heading(doc, '7 商业价值与应用前景', 1)
    para(doc, '充储场站与充电网络规模扩张，告警处置依赖经验且耗时。本作品将「证据→候选→补证→人工」的诊断流程标准化：一次诊断输出可核验的候选与补证清单，减少漏判与误判；Agent 与本地合同可跨场站复用；与工单、大屏联动后可形成告警处置闭环。满足大赛对万智平台的使用要求：已发布 Agent、完成与实际业务直接相关的智能体任务（真实 API 证据诊断），并提供公开 URL 供评审访问。', indent_chars=2)

    # 8 边界声明
    heading(doc, '8 边界声明', 1)
    para(doc, '演示输入为手工构造的合成样例（SYNTHETIC_DEVELOPMENT_FIXTURE）与参照典型充储场站参数构造的行业演示数据（INDUSTRY_DEMO_FIXTURE），均非真实场站遥测；E04 电压为本地声明 fixture 的模拟读数，输出中已由模型标注。作品真实 API 调用范围限于只读证据诊断，未部署任何设备控制能力。', indent_chars=2)
    para(doc, '参考资料：云桥杯大赛官网（https://yunqiao.jpx.com/cup/）、万智开放 API（https://www.wanzhi.com/app/api/v1/open-api）、仓库 docs/integrations/wanzhi/README.md 与 observations/*.md、decision.schema.json。', indent_chars=2)

    doc.save(OUT)
    print('SAVED', OUT)


if __name__ == '__main__':
    main()
